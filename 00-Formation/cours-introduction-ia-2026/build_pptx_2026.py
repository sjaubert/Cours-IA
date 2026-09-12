"""
Script de génération du diaporama actualisé 2026 :
00-Formation/cours-introduction-ia-2026/Diaporama_COURS_Introduction_IA_2026.pptx
Basé sur l'original : 00-Formation/Module Formation/Diaporama COURS Introduction IA.pptx
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(os.path.dirname(SCRIPT_DIR))
ORIGINAL_PPTX = os.path.join(PROJECT_ROOT, "00-Formation", "Module Formation", "Diaporama COURS Introduction IA.pptx")
OUTPUT_PPTX = os.path.join(SCRIPT_DIR, "Diaporama_COURS_Introduction_IA_2026.pptx")
ASSETS_DIR = os.path.join(SCRIPT_DIR, "assets")
LOGO_PATH = os.path.join(PROJECT_ROOT, "logo_uimm_placeholder.jpg")

# Palette graphique industrielle moderne (Dark Slate & Cyan/Indigo)
COLOR_BG = RGBColor(11, 17, 32)         # #0B1120
COLOR_CARD_BG = RGBColor(30, 41, 59)    # #1E293B
COLOR_IMG_BG = RGBColor(15, 23, 42)     # #0F172A
COLOR_BORDER = RGBColor(51, 65, 85)     # #334155
COLOR_TITLE = RGBColor(248, 250, 252)   # #F8FAFC
COLOR_SUBTITLE = RGBColor(56, 189, 248) # #38BDF8 (Cyan)
COLOR_TEXT = RGBColor(226, 232, 240)    # #E2E8F0
COLOR_MUTED = RGBColor(148, 163, 184)   # #94A3B8
COLOR_ACCENT = RGBColor(129, 140, 248)  # #818CF8 (Indigo)
COLOR_LINK = RGBColor(96, 165, 250)     # #60A5FA (Light Blue)

NEW_SLIDES_DATA = [
    {
        "module": "MODULE 3 — MUTATIONS GÉNÉRATIVES & AGENTS (2022-2026)",
        "title": "Révolution Générative & Mécanisme d'Attention des Transformers",
        "subtitle": "De la classification discriminante à la génération contextuelle",
        "image": "transformer_architecture.png",
        "caption": "Architecture Encodeur-Décodeur Transformer (Vaswani et al., NeurIPS 2017)",
        "bullets": [
            "Rupture avec les RNN/LSTM : Abandon du traitement séquentiel au profit d'une parallélisation massive sur GPU grâce à l'auto-attention.",
            "Formulation mathématique fondamentale : Attention(Q, K, V) = softmax(Q * K^T / sqrt(d_k)) * V calculant le poids contextuel de chaque token.",
            "Changement de paradigme : Les modèles ne classent plus de simples données, ils prédisent le prochain token dans un espace latent ultra-dimensionnel.",
            "Impacts atelier : Compréhension instantanée de documentations techniques, génération de code automate et synthèse de retours d'expérience."
        ],
        "sources": [
            ("Vaswani et al. (2017)", "Attention Is All You Need (arXiv:1706.03762)", "https://arxiv.org/abs/1706.03762"),
            ("Stanford HAI", "Artificial Intelligence Index Report 2024", "https://aiindex.stanford.edu/report/")
        ]
    },
    {
        "module": "MODULE 3 — MUTATIONS GÉNÉRATIVES & AGENTS (2022-2026)",
        "title": "L'IA Multimodale en Atelier : Vision, Son, Capteurs & Télémétrie",
        "subtitle": "Fusion de données hétérogènes dans un même espace d'encodage",
        "image": "multimodal_industrial_ai.png",
        "caption": "Pipeline d'ingestion multimodale et diagnostic atelier unifié",
        "bullets": [
            "Espace latent partagé : Alignement mathématique des représentations visuelles (pixels), sonores (spectrogrammes) et textuelles (rapports).",
            "Contrôle qualité haute cadence : Segmentation temps réel de défauts de surface sur pièces usinées avec localisation millimétrique.",
            "Maintenance prédictive acoustique : Corrélation instantanée entre vibrations anormales d'un roulement et historique de maintenance GMAO.",
            "Assistance opérateur terrain : Prise de photo d'un code défaut sur automate avec diagnostic instantané et procédure de consignation validée."
        ],
        "sources": [
            ("Meta AI Research", "Segment Anything Model (arXiv:2304.02643)", "https://arxiv.org/abs/2304.02643"),
            ("Journal of Manufacturing Systems", "Multimodal Industrial Inspection (2024)", "https://doi.org/10.1016/j.jmsy.2024.01.008")
        ]
    },
    {
        "module": "MODULE 3 — MUTATIONS GÉNÉRATIVES & AGENTS (2022-2026)",
        "title": "Modèles de Raisonnement & Chaîne de Pensée (System 1 vs System 2)",
        "subtitle": "Passage de la réponse réflexe au calcul logique pas-à-pas",
        "image": "reasoning_chain_of_thought.png",
        "caption": "Comparaison cognitive : Génération token réflexe vs Délibération pas-à-pas avec CoT",
        "bullets": [
            "Limites du Système 1 (LLM standards) : Réponse statistique token par token sans retour en arrière, risque élevé d'hallucinations sur calculs complexes.",
            "Paradigme Système 2 (o1, DeepSeek-R1) : Génération préalable de tokens de délibération cachés (Chain-of-Thought) avant d'émettre la conclusion.",
            "Exploration et retour arrière (backtracking) : L'IA formule des hypothèses, teste leur cohérence logique, backtracke si nécessaire et s'auto-corrige.",
            "Applications industrielles : Ordonnancement complexe de production sous contraintes multiples, calcul de tolérancement et diagnostics sécurité."
        ],
        "sources": [
            ("Wei et al. (Google Brain)", "Chain-of-Thought Prompting (arXiv:2201.11903)", "https://arxiv.org/abs/2201.11903"),
            ("DeepSeek-AI (2025)", "DeepSeek-R1: Incentivizing Reasoning in LLMs", "https://arxiv.org/abs/2501.12948")
        ]
    },
    {
        "module": "MODULE 3 — MUTATIONS GÉNÉRATIVES & AGENTS (2022-2026)",
        "title": "Architecture RAG Industrielle (Retrieval-Augmented Generation)",
        "subtitle": "Ancrage des modèles sur la documentation privée de l'usine",
        "image": "rag_architecture.png",
        "caption": "Architecture de génération augmentée par récupération vectorielle d'atelier",
        "bullets": [
            "Le problème fondamental : Un modèle pré-entraîné ignore vos plans internes, vos gammes d'usinage et vos procédures confidentielles.",
            "Principe du RAG : Découpage des documents techniques en fragments (chunks), indexation vectorielle (embeddings) et recherche de similarité cosinus.",
            "Injection contextuelle : Les extraits exacts des manuels machines sont injectés dans le prompt en temps réel pour guider la réponse du LLM.",
            "Traçabilité & Auditabilité : Chaque recommandation générée mentionne explicitement le numéro de page et le document technique source vérifié."
        ],
        "sources": [
            ("Lewis et al. / Meta AI", "Retrieval-Augmented Generation (arXiv:2005.11401)", "https://arxiv.org/abs/2005.11401"),
            ("Microsoft Architecture Center", "Enterprise RAG Design Patterns (2024)", "https://techcommunity.microsoft.com/")
        ]
    },
    {
        "module": "MODULE 3 — MUTATIONS GÉNÉRATIVES & AGENTS (2022-2026)",
        "title": "Systèmes Agentiques & Protocole Ouvert MCP (Model Context Protocol)",
        "subtitle": "De l'assistant conversationnel à l'agent d'action interconnecté",
        "image": "mcp_protocol_architecture.png",
        "caption": "Standard ouvert MCP : Interconnexion unifiée entre LLM et systèmes industriels (Anthropic, 2024)",
        "bullets": [
            "Du Chatbot à l'Agent : Capacité autonome à percevoir l'environnement, planifier des sous-tâches, appeler des outils (Tool Calling) et vérifier le résultat.",
            "Le standard MCP (Anthropic fin 2024) : Protocole ouvert évitant le développement de connecteurs propriétaires ad-hoc pour chaque outil d'usine.",
            "Passerelles d'atelier : Serveurs MCP interfaçant directement l'ERP (SAP), la GMAO, les bases SQL et les automates industriels via OPC-UA.",
            "Gouvernance & Sécurité : Cloisonnement strict des accès, audit des traces d'exécution et nécessité impérative du contrôle humain (Human-in-the-Loop)."
        ],
        "sources": [
            ("Anthropic", "Model Context Protocol Specification (2024)", "https://modelcontextprotocol.io/"),
            ("Linux Foundation", "Open Standards for Agentic AI Systems (2025)", "https://www.linuxfoundation.org/")
        ]
    },
    {
        "module": "MODULE 4 — SOUVERAINETÉ, RÉGULATION & MÉTHODOLOGIE (2024-2026)",
        "title": "Régulation Européenne : Le Règlement AI Act (2024-2026)",
        "subtitle": "Cadre juridique contraignant et classification par niveau de risque",
        "image": "ai_act_risk_pyramid.png",
        "caption": "Pyramide des 4 niveaux de risque du Règlement (UE) 2024/1689 (AI Act)",
        "bullets": [
            "Réglementation pionnière : Règlement (UE) 2024/1689 adopté le 13 juin 2024, avec sanctions financières pouvant atteindre 35 millions d'euros ou 7% du CA.",
            "Approche fondée sur le risque : 4 catégories de risques allant de l'interdit pur au risque minimal, avec exigences proportionnées à l'impact humain.",
            "Impacts systèmes industriels : Les IA intégrées aux machines-outils de sécurité, à la robotique collaborative et au tri RH sont classées 'Haut Risque'.",
            "Obligations constructeurs : Analyse de risques continue, gouvernance des données d'entraînement, robustesse cyber, traçabilité et marquage CE obligatoire."
        ],
        "sources": [
            ("Journal Officiel de l'UE", "Règlement (UE) 2024/1689 du 13 juin 2024", "https://eur-lex.europa.eu/eli/reg/2024/1689/oj"),
            ("Commission Européenne", "Bureau Européen de l'IA (AI Office)", "https://digital-strategy.ec.europa.eu/en/policies/ai-office")
        ]
    },
    {
        "module": "MODULE 4 — SOUVERAINETÉ, RÉGULATION & MÉTHODOLOGIE (2024-2026)",
        "title": "Souveraineté Numérique & Modèles Ouverts (Open-Weights)",
        "subtitle": "Maîtrise de la propriété intellectuelle et indépendance technologique",
        "image": "sovereign_open_weights_models.png",
        "caption": "Écosystème des modèles ouverts (Mistral AI, Meta Llama) vs Modèles propriétaires Cloud",
        "bullets": [
            "Vulnérabilités du Cloud propriétaire : Risque d'extraterritorialité juridique (Cloud Act), fuite potentielle de secrets de fabrication et dépendance tarifaire.",
            "L'alternative Open-Weights : Poids de modèles librement téléchargeables (Mistral, Llama, Qwen) offrant des performances comparables aux modèles fermés.",
            "Déploiement On-Premise Air-Gapped : Hébergement local sur serveurs internes à l'usine, totalement déconnectés d'Internet (étanchéité des données).",
            "Personnalisation et souveraineté : Fine-tuning sans restriction sur les lexiques métiers internes, coût marginal par requête nul en phase de croisière."
        ],
        "sources": [
            ("Mistral AI", "Modèles Ouverts Européens & Architecture Mixtral", "https://mistral.ai/news/"),
            ("ANSSI", "Recommandations de sécurité pour l'IA générative (2024)", "https://www.ssi.gouv.fr/")
        ]
    },
    {
        "module": "MODULE 4 — SOUVERAINETÉ, RÉGULATION & MÉTHODOLOGIE (2024-2026)",
        "title": "Déploiement Local & Frugalité : SLM et Quantification Numérique",
        "subtitle": "Rendre l'IA embarquable sur matériel industriel standard",
        "image": "frugal_ai_comparison.png",
        "caption": "Comparatif d'empreinte VRAM : Impact de la quantification 4-bit / 8-bit sur SLM et LLM",
        "bullets": [
            "L'essor des Petits Modèles (SLM - 1B à 8B) : Modélisation optimisée sur données d'entraînement filtrées, rivalisant avec les anciens mastodontes.",
            "Mécanisme de Quantification (AWQ, GGUF, GPTQ) : Réduction de précision des poids de FP16 (16 bits) à INT4 (4 bits) avec perte de fidélité inférieure à 1%.",
            "Diviser l'empreinte mémoire par 4 : Un modèle 7B en FP16 nécessite 14 à 16 Go de VRAM ; quantifié en 4-bit, il s'exécute sur seulement 4.5 Go de VRAM.",
            "Embarquement Edge usine : Exécution fluide sur cartes d'automatisme (Nvidia Jetson, NPU) ou PC d'atelier d'entrée de gamme sans abonnement cloud."
        ],
        "sources": [
            ("Dettmers et al. (Washington Univ.)", "QLoRA: Efficient Finetuning (arXiv:2305.14314)", "https://arxiv.org/abs/2305.14314"),
            ("Microsoft Research", "The Phi-3 Technical Report (arXiv:2404.14219)", "https://arxiv.org/abs/2404.14219")
        ]
    },
    {
        "module": "MODULE 4 — SOUVERAINETÉ, RÉGULATION & MÉTHODOLOGIE (2024-2026)",
        "title": "Enjeux Énergétiques & Bilan Carbone de l'IA Industrielle",
        "subtitle": "Mesurer et optimiser le coût électrique du cycle de vie des modèles",
        "image": "energy_footprint_ai.png",
        "caption": "Consommation électrique par requête (Wh) et répartition du cycle de vie IA (IEA / Luccioni 2024)",
        "bullets": [
            "Dominance de l'inférence : Contrairement aux idées reçues, 75% à 85% de l'empreinte énergétique d'une IA provient de son usage quotidien en production.",
            "Échelle de consommation : Requête web classique (~0.3 Wh), LLM 7B local quantifié (~1.2 Wh), modèle Cloud 70B (~4.5 Wh), modèle de raisonnement complexe (~18 Wh).",
            "Principe de frugalité : Adapter rigoureusement la taille du modèle à la complexité de la tâche (ne pas allouer un LLM 70B à une simple classification d'alertes).",
            "Indicateurs RSE & Usine Verte : Prise en compte du PUE (Power Usage Effectiveness) des baies serveurs et du coût en kWh par pièce contrôlée."
        ],
        "sources": [
            ("Luccioni et al. (Hugging Face)", "Power Hungry Processing (arXiv:2311.16863)", "https://arxiv.org/abs/2311.16863"),
            ("Agence Internationale de l'Énergie (IEA)", "Electricity Report 2024: Data Centres & AI", "https://www.iea.org/reports/electricity-2024")
        ]
    },
    {
        "module": "MODULE 4 — SOUVERAINETÉ, RÉGULATION & MÉTHODOLOGIE (2024-2026)",
        "title": "Éthique & Métacognition : Les 4 Pièges Cognitifs de l'Ingénieur",
        "subtitle": "Penser avec la machine sans renoncer à l'esprit critique",
        "image": "metacognition_biases_matrix.png",
        "caption": "Matrice des 4 pièges cognitifs face à l'IA et contre-mesures méthodologiques",
        "bullets": [
            "Biais d'automatisation : Tendance psychologique à déléguer aveuglément la décision à l'algorithme sans vérification contradictoire sur pièce réelle.",
            "Illusion de compétence : L'éloquence du langage d'un modèle génère une fausse impression d'omniscience technique, masquant les failles physiques.",
            "Biais de confirmation renforcé : Risque d'utiliser l'IA pour conforter une intuition erronée sans tester les conditions aux limites de sécurité.",
            "Atrophie des compétences : Perte progressive de savoir-faire critique si les analyses de panne fondamentales ne sont plus pratiquées manuellement.",
            "Posture de l'Ingénieur : 'Penser contre soi-même' — imposition stricte d'un protocole de validation indépendante (Human-in-the-Loop)."
        ],
        "sources": [
            ("Parasuraman & Manzey", "Complacency and Bias in Human Automation (Human Factors 2010)", "https://doi.org/10.1177/0018720810376055"),
            ("AFNOR Spec IA (2024)", "Guide de fiabilité des systèmes à base d'IA", "https://www.afnor.org/")
        ]
    },
    {
        "module": "MODULE 4 — SOUVERAINETÉ, RÉGULATION & MÉTHODOLOGIE (2024-2026)",
        "title": "Synthèse & Feuille de Route Opérationnelle de l'Ingénieur 2026",
        "subtitle": "Cycle méthodologique en 4 étapes pour industrialiser un projet d'IA",
        "image": "industrial_ai_roadmap_2026.png",
        "caption": "Feuille de route en 4 étapes : Du cadrage ROI à la gouvernance conforme AI Act",
        "bullets": [
            "Étape 1 - Cadrage & ROI : Qualification du besoin métier, calcul de retour sur investissement industriel et classification préliminaire de risque AI Act.",
            "Étape 2 - Architecture Souveraine : Sélection du modèle (SLM ou LLM spécialisé), choix du runtime local (Ollama / vLLM) et isolation réseau usine.",
            "Étape 3 - Intégration RAG & MCP : Raccordement sécurisé à la documentation technique, protocoles d'automatisme (OPC-UA) et supervision agentique.",
            "Étape 4 - Gouvernance & Cycle de Vie : Documentation technique obligatoire, surveillance de la dérive du modèle (Drift), formation des équipes et audit continu."
        ],
        "sources": [
            ("Organisation Internationale de Normalisation", "ISO/IEC 42001:2023 - Management de l'IA", "https://www.iso.org/standard/81230.html"),
            ("Ministère de l'Économie et des Finances", "Stratégie Nationale pour l'IA Industrielle", "https://www.entreprises.gouv.fr/")
        ]
    }
]


def update_slide_1(slide):
    """Actualise la diapositive de couverture (titre, date et mentions UIMM)."""
    # 1. Update text shapes
    for sh in slide.shapes:
        if sh.has_text_frame:
            tf = sh.text_frame
            for p in tf.paragraphs:
                if "23/12/2025" in p.text:
                    p.text = "Septembre 2026"
                    p.font.size = Pt(14)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(30, 41, 59)
                elif "Bases, concepts et histoire" in p.text:
                    p.text = "Bases, concepts, mutations génératives et souveraineté industrielle"
                    p.font.size = Pt(18)
                    p.font.color.rgb = RGBColor(51, 65, 85)

    # 2. Add institution & trainer badge on slide 1 (top left)
    badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10.0), Inches(0.8))
    btf = badge_box.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.text = "Pôle Formation UIMM - CVDL  |  Formateur : S. JAUBERT"
    bp.font.name = "Calibri"
    bp.font.size = Pt(14)
    bp.font.bold = True
    bp.font.color.rgb = RGBColor(15, 23, 42)


def update_references_slide(slide, total_slides):
    """Réorganise la diapositive finale de références en deux colonnes équilibrées."""
    # Supprimer l'ancienne boîte de texte qui déborde
    shapes_to_remove = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top > Inches(1.5):
            shapes_to_remove.append(sh)
    for sh in shapes_to_remove:
        sp_elem = sh._element
        sp_elem.getparent().remove(sp_elem)

    # Colonne 1 : Références Fondamentales & Pédagogiques
    col1_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    c1_tf = col1_box.text_frame
    c1_tf.word_wrap = True

    h1 = c1_tf.paragraphs[0]
    h1.text = "Références Fondamentales & Générales"
    h1.font.name = "Calibri"
    h1.font.size = Pt(13)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(15, 23, 42)
    h1.space_after = Pt(8)

    refs_fondamentales = [
        "Conférence Yann LeCun (Collège de France) : Atteindre l'intelligence humaine",
        "Stéphane Roder (Eyrolles) : Guide pratique de l'IA dans l'entreprise",
        "CNRS / FIDLE : Introduction approfondie au Deep Learning",
        "Collège de France : Colloque L'IA et ses défis scientifiques et sociétaux",
        "Université de Genève & Inspé : Guides pratiques d'accompagnement à l'IA",
        "DGE / Ministère : Décryptage des opportunités de l'IA pour les entreprises"
    ]
    for r in refs_fondamentales:
        p = c1_tf.add_paragraph()
        p.text = f"• {r}"
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(6)

    # Colonne 2 : Références Réglementaires & Techniques 2024-2026
    col2_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.6), Inches(5.2))
    c2_tf = col2_box.text_frame
    c2_tf.word_wrap = True

    h2 = c2_tf.paragraphs[0]
    h2.text = "Références Techniques & Normes 2024-2026"
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(2, 132, 199)
    h2.space_after = Pt(8)

    refs_modernes = [
        "Règlement (UE) 2024/1689 (AI Act) : Cadre légal et exigences des systèmes à haut risque",
        "ISO/IEC 42001:2023 : Système de management de l'intelligence artificielle (SMIA)",
        "Vaswani et al. (2017) : 'Attention Is All You Need' — Mécanisme Transformer fondateur",
        "Anthropic (2024) : Spécification du standard ouvert MCP (Model Context Protocol)",
        "AFNOR Spec IA (2024) : Bonnes pratiques d'ingénierie et de robustesse des modèles",
        "ANSSI (2024) : Guide de sécurité pour l'IA générative et souveraineté des données",
        "Luccioni et al. / AIE (2024) : Bilan énergétique et mesure de l'empreinte carbone de l'IA"
    ]
    for r in refs_modernes:
        p = c2_tf.add_paragraph()
        p.text = f"• {r}"
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(6)


def add_footer_to_slide(slide, slide_num, total_slides):
    """Ajoute le bandeau de conformité UIMM en pied de page sur chaque diapositive."""
    ft_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
    tf = ft_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"Pôle Formation UIMM - CVDL  |  Formateur : S. JAUBERT  |  Formation Introduction à l'IA (2026)  |  {slide_num}/{total_slides}"
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = RGBColor(148, 163, 184)


def create_modern_slide(prs, data):
    """Crée une diapositive moderne pour les modules 3 et 4 avec richesse visuelle et sources web."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Background fill sombre élégant (#0B1120)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

    # 2. Header : Badge Module
    badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(9.2), Inches(0.3))
    btf = badge_box.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.text = data["module"]
    bp.font.name = "Calibri"
    bp.font.size = Pt(10.5)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_SUBTITLE

    # 3. Header : Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(9.2), Inches(0.75))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.text = data["title"]
    tp.font.name = "Calibri"
    tp.font.size = Pt(19)
    tp.font.bold = True
    tp.font.color.rgb = COLOR_TITLE

    # 4. Header : Branding UIMM & Logo (Top Right)
    brand_box = slide.shapes.add_textbox(Inches(10.1), Inches(0.35), Inches(2.2), Inches(0.55))
    brtf = brand_box.text_frame
    brtf.word_wrap = True
    brp1 = brtf.paragraphs[0]
    brp1.text = "Pôle Formation UIMM - CVDL"
    brp1.font.name = "Calibri"
    brp1.font.size = Pt(8.5)
    brp1.font.bold = True
    brp1.font.color.rgb = COLOR_TITLE
    brp1.alignment = PP_ALIGN.RIGHT

    brp2 = brtf.add_paragraph()
    brp2.text = "Formateur : S. JAUBERT"
    brp2.font.name = "Calibri"
    brp2.font.size = Pt(8)
    brp2.font.color.rgb = COLOR_MUTED
    brp2.alignment = PP_ALIGN.RIGHT

    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(12.45), Inches(0.32), width=Inches(0.55), height=Inches(0.58))

    # 5. Left Column : Card Container
    card_left = Inches(0.8)
    card_top = Inches(1.45)
    card_w = Inches(5.8)
    card_h = Inches(5.4)

    rect_left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, card_left, card_top, card_w, card_h)
    rect_left.fill.solid()
    rect_left.fill.fore_color.rgb = COLOR_CARD_BG
    rect_left.line.color.rgb = COLOR_BORDER
    rect_left.line.width = Pt(1.2)

    # Subtitle inside Left Card
    sub_box = slide.shapes.add_textbox(Inches(0.95), Inches(1.55), Inches(5.5), Inches(0.45))
    sft = sub_box.text_frame
    sft.word_wrap = True
    sp = sft.paragraphs[0]
    sp.text = data["subtitle"]
    sp.font.name = "Calibri"
    sp.font.size = Pt(12)
    sp.font.bold = True
    sp.font.color.rgb = COLOR_SUBTITLE

    # Bullets inside Left Card
    bullets_box = slide.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.5), Inches(3.2))
    btf = bullets_box.text_frame
    btf.word_wrap = True
    for idx, bullet in enumerate(data["bullets"]):
        p = btf.paragraphs[0] if idx == 0 else btf.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.name = "Calibri"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(7)

    # Sources block inside Left Card (bottom)
    sources_box = slide.shapes.add_textbox(Inches(0.95), Inches(5.35), Inches(5.5), Inches(1.4))
    stf = sources_box.text_frame
    stf.word_wrap = True
    shp = stf.paragraphs[0]
    shp.text = "Sources & Références Officielles :"
    shp.font.name = "Calibri"
    shp.font.size = Pt(9)
    shp.font.bold = True
    shp.font.color.rgb = COLOR_MUTED
    shp.space_after = Pt(2)

    for author, title, url in data["sources"]:
        sp = stf.add_paragraph()
        sp.text = f"› {author} : {title} — {url}"
        sp.font.name = "Calibri"
        sp.font.size = Pt(8)
        sp.font.color.rgb = COLOR_LINK
        sp.space_after = Pt(1)

    # 6. Right Column : Card Container for Diagram
    img_card_left = Inches(6.8)
    img_card_top = Inches(1.45)
    img_card_w = Inches(5.8)
    img_card_h = Inches(5.4)

    rect_right = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, img_card_left, img_card_top, img_card_w, img_card_h)
    rect_right.fill.solid()
    rect_right.fill.fore_color.rgb = COLOR_IMG_BG
    rect_right.line.color.rgb = COLOR_BORDER
    rect_right.line.width = Pt(1.2)

    # Place Image inside Right Card with preserved aspect ratio
    image_file = os.path.join(ASSETS_DIR, data["image"])
    if os.path.exists(image_file):
        with Image.open(image_file) as im:
            im_w, im_h = im.size
            im_aspect = im_w / im_h

        # Available space for picture
        avail_w = Inches(5.5)
        avail_h = Inches(4.5)
        avail_aspect = 5.5 / 4.5

        if im_aspect > avail_aspect:
            # Width constrained
            pic_w = avail_w
            pic_h = Inches(5.5 / im_aspect)
        else:
            # Height constrained
            pic_h = avail_h
            pic_w = Inches(4.5 * im_aspect)

        pic_left = img_card_left + (img_card_w - pic_w) / 2
        pic_top = img_card_top + Inches(0.2) + (avail_h - pic_h) / 2

        slide.shapes.add_picture(image_file, pic_left, pic_top, width=pic_w, height=pic_h)

    # Caption Box below image
    caption_box = slide.shapes.add_textbox(Inches(6.9), Inches(6.35), Inches(5.6), Inches(0.45))
    ctf = caption_box.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = data["caption"]
    cp.font.name = "Calibri"
    cp.font.size = Pt(8.5)
    cp.font.italic = True
    cp.font.color.rgb = COLOR_MUTED
    cp.alignment = PP_ALIGN.CENTER

    return slide


def main():
    print("=== Démarrage de la génération du diaporama 2026 ===")
    print(f"Original PPTX : {ORIGINAL_PPTX}")
    print(f"Output PPTX   : {OUTPUT_PPTX}")

    if not os.path.exists(ORIGINAL_PPTX):
        raise FileNotFoundError(f"Fichier introuvable : {ORIGINAL_PPTX}")

    prs = pptx.Presentation(ORIGINAL_PPTX)
    orig_count = len(prs.slides)
    print(f"Nombre de diapositives d'origine : {orig_count}")

    # Total slides = 66 originales + 11 nouvelles = 77
    total_slides = orig_count + len(NEW_SLIDES_DATA)

    # 1. Mise à jour de la Diapositive 1 (Couverture)
    print("Mise à jour de la diapositive 1...")
    update_slide_1(prs.slides[0])

    # 2. Création et insertion des 11 nouvelles diapositives
    # Dans l'original :
    # Indices 0..63 = slides 1..64 (Technique ML/DL)
    # Index 64 = slide 65 (Citation philosophique de S. Jaubert)
    # Index 65 = slide 66 (Références)
    # On insère les 11 slides avant la slide 65 (donc à l'index 64)
    print("Génération et insertion des 11 nouvelles diapositives illustrées...")
    sldIdLst = prs.slides._sldIdLst

    for k, data in enumerate(NEW_SLIDES_DATA):
        print(f"  -> Ajout slide {65 + k}/{total_slides} : {data['title'][:45]}...")
        create_modern_slide(prs, data)
        # Déplacer l'élément ajouté en fin de liste vers l'index 64 + k
        new_elem = sldIdLst[-1]
        sldIdLst.remove(new_elem)
        sldIdLst.insert(64 + k, new_elem)

    # 3. Actualisation de la slide de références (désormais la dernière slide, index 76 / slide 77)
    print("Mise en page à deux colonnes de la diapositive de références...")
    ref_slide = prs.slides[76]
    update_references_slide(ref_slide, total_slides)

    # 4. Ajout propre des footers numérotés sur TOUTES les diapositives de 2 à 77
    print("Application du bandeau de conformité UIMM sur l'ensemble des diapositives 2 à 77...")
    for idx in range(1, len(prs.slides)):
        add_footer_to_slide(prs.slides[idx], idx + 1, total_slides)

    # 5. Sauvegarde dans le nouveau fichier
    print(f"Sauvegarde du nouveau diaporama dans {OUTPUT_PPTX}...")
    os.makedirs(os.path.dirname(OUTPUT_PPTX), exist_ok=True)
    prs.save(OUTPUT_PPTX)

    print(f"=== Génération terminée avec succès ! ===")
    print(f"Ancien fichier préservé : {ORIGINAL_PPTX}")
    print(f"Nouveau fichier créé    : {OUTPUT_PPTX}")
    print(f"Nombre total de diapositives : {len(prs.slides)}")


if __name__ == "__main__":
    main()
